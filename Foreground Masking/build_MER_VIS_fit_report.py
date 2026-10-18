import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

root=Path('MER_VIS_report_work');s=json.loads((root/'statistics.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(24);doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(12)
footer=sec.footer.paragraphs[0];footer.text='Euclid MER VIS geometry assessment  •  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
def p(text):doc.add_paragraph(text)
def h(text):doc.add_heading(text,2)
def page(text):doc.add_page_break();doc.add_heading(text,1)
def bullet(text):doc.add_paragraph(text,style='List Bullet')
def pic(name,caption):
 doc.add_picture(str(root/name),width=Inches(6.85))
 a=doc.add_paragraph(caption)
 for r in a.runs:r.font.size=Pt(9)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
 for i,row in enumerate(t.rows):
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EEF3' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));tcpr.append(sh)
   margins=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');margins.append(x)
   tcpr.append(margins)
   for a in cell.paragraphs:
    a.paragraph_format.space_after=Pt(2)
    for r in a.runs:r.font.size=Pt(9.5);r.bold=i==0
 borders=OxmlElement('w:tblBorders')
 for side in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
 t._tbl.tblPr.append(borders)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
def pct(n,d=2777):return f'{100*n/d:.1f}%'

doc.add_paragraph('Euclid MER VIS galaxy geometry fitting report',style='Title')
p('Assessment of the completed local run and comparison with public Euclid Q1 morphology. Prepared 18 October 2026 for the 2,777 candidate barred spiral galaxies selected from Zenodo record 15106473.')
h('Principal findings')
p('All 2,777 targets were attempted. The algorithm returned 2,207 provisional geometry estimates, of which 546 carry no automatic quality flags. Those 546 are a review shortlist, not 546 independently validated disc and bar measurements. The remaining fitted objects and failures require additional checks or revised fitting ranges.')
table(['Outcome','Galaxies','Share of all targets'],[
 ['Provisional fit with no quality flags',546,pct(546)],['Provisional fit with one or more flags',1661,pct(1661)],['Fit failed',570,pct(570)],['Total attempted',2777,'100.0%']], [4.2,1,1.6])
p('A material issue was found in the catalogue PA conversion: the stored Euclid sky-angle comparison is offset by approximately a quarter turn for much of the sample. Treat the affected PA comparison columns as unvalidated until the convention is resolved. This does not establish a failure of the independent Photutils fits and does not alter the inclination comparisons.')
pic('outcomes_compact.png','Figure 1. Outcome counts and inclination distributions. The unflagged subset is selected by the fitter’s thresholds and is biased towards more inclined systems.')
p('The median calculated inclination is 43.8° across all fits and 53.2° in the unflagged subset. Inclinations assume an infinitely thin circular disc, q0 = 0. The observed bar and disc PAs are axial directions in degrees east of north, modulo 180°.')
p('No independent Euclid bar PA is available in the queried Q1 morphology tables. All separate disc-fit axis ratios are unpopulated for this sample. Euclid detection and single-Sérsic shapes therefore provide whole-object comparison proxies rather than disc or bar ground truth.')

page('Data and fitting method')
h('Sample and source products')
p('The input comprises 2,777 unique object IDs and native galaxy-centred MER VIS background-subtracted science cutouts, drawn from 309 selected mosaics. The first-25 limit used in the initial exploratory work is absent from this full selection. The catalogue criteria require Q1_R1, a false training-cut warning, and fractions ≥0.8 for featured/disc, not edge-on, spiral arms and strong plus weak bar. These conditional morphology fractions select candidates; they do not prove each galaxy has a bar.')
p('All 2,777 IDs matched both cached Euclid catalogue tables. Catalogue values were retrieved on 16 October 2026 and are not a new live archive query performed for this report. The fitting run completed on 18 October 2026. The main run summary records 6,229 seconds, approximately 104 minutes; calculation timestamps span 08:19:44 to 10:06:07 UTC, including pilot and interruption gaps. Median per-object recorded computation was 3.74 seconds; the 90th percentile was 7.70 seconds.')
h('Image based calculation')
p('The existing Photutils ellipse algorithm was reused locally with two independent workers. The catalogue centre is fixed. Fitting uses 2 × 2 averaged pixels, a sigma-clipped local background estimate, masks for disconnected neighbours and clipped ellipse samples. The native input images remain unchanged. These science-only cutouts have no FLG map, so detector flags cannot be applied; no RMS weighting or PSF-convolved component model is used.')
p('Multiple starting ellipses are tried. Reliable samples require a valid fit, stop code 0, ≥70% sample coverage, positive intensity and intensity above five times its formal error. At least ten reliable ellipses overall and three in each selected bar and disc range are required. Automatic bar selection follows an inner ellipticity peak; a separate outer radial interval supplies the disc shape. Spiral arms, bulges, rings and connected companions can affect either interval.')
p('Disc b/a is one minus the median outer-range ellipticity, and inclination is arccos(b/a). Bar PA and disc PA are axial circular means; image directions are converted through the native WCS. PA scatter describes radial variability, not calibrated uncertainty on the adopted geometry. The inner ellipticity maximum and its radial boundary are not a rigorously measured bar end.')
h('Meaning of the catalogue columns')
table(['Measurement','Meaning and comparison limit'],[
 ['POSITION_ANGLE and SEMIMAJOR_AXIS','Detection shape direction and moment size. They are not a separately fitted bar PA or bar length.'],
 ['ELLIPTICITY','Detection shape: use q = 1 − ellipticity for a thin-disc inclination proxy.'],
 ['Single Sérsic VIS axis ratio','Whole-galaxy fitted b/a, available for all 2,777 objects; not an isolated outer disc.'],
 ['Separate disc fit and bar PA','Disc-fit b/a is empty for this sample; no independent bar angle is provided.']], [2.25,4.55])
p('References: the Euclid MER Morphology Cookbook, Final Catalog definitions and SExtractor shape definitions are listed on the final page. The PA interpretation conflict is examined separately below.')

page('Failures quality flags and review priorities')
h('Why 570 fits failed')
table(['Failure category','Count','Share of failures'],[
 ['Too few samples in selected bar or disc ranges',431,pct(431,570)],
 ['Fewer than ten reliable isophotes overall',138,pct(138,570)],
 ['Catalogue centre not on a detected component',1,pct(1,570)]], [4.4,.7,1.7])
p('These are recorded fitting/selection failures, not evidence that the objects are not barred spirals. The 431 radial-selection failures are the largest recovery opportunity: inspect their saved profiles and choose physically appropriate ranges. The 138 short profiles include 64 with zero reliable isophotes. Masks, low gradients, neighbouring light, centre placement and limited coverage are possible causes; the error strings do not diagnose which cause applies to each object.')
h('Flags among the 2207 provisional fits')
flagrows=[['Bar range upper bound below 1.5 arcsec',1210,pct(1210,2207)],['Disc PA radial scatter above 10°',791,pct(791,2207)],['Disc ellipticity span above 0.15',573,pct(573,2207)],['Nearly face-on axis ratio above 0.9',200,pct(200,2207)],['Bar PA radial scatter above 10°',28,pct(28,2207)],['Median bar ellipticity below 0.25',7,pct(7,2207)]]
table(['Quality condition','Count','Share of fits'],flagrows,[4.4,.7,1.7])
p('Flags overlap and must not be added to obtain a galaxy count. There are 1,661 distinct flagged fits. No high-inclination flag at ≥80° was recorded. A clear flag list only means these numerical thresholds were passed; it does not establish the correct physical disc or bar region.')
h('Coverage and selection effects')
p('There are 62 PARTIAL cutouts overall, including 51 provisional fits and 19 unflagged fits. PARTIAL denotes a requested sky box crossing a mosaic edge; the galaxy can still be complete. Inspect these images rather than automatically rejecting them. A conservative first review list could start with the 527 unflagged fits whose requested boxes are not partial, while retaining the other 19 for coverage inspection.')
p('The unflagged subset has median inclination 53.2°, compared with 43.8° for all fits. This difference is expected from geometry-dependent measurement difficulty and the nearly face-on flag; it means the 546-object subset is not an unbiased scientific subsample. Small-bar flags also strongly affect which galaxies pass.')

page('Position angle comparison requires a convention audit')
p('The stored euclid_detection_pa_sky_deg was created by treating POSITION_ANGLE as SExtractor THETA_IMAGE and converting that pixel direction through the cutout WCS. TAP metadata describes the column as counterclockwise from x. The morphology cookbook instead describes an NE–SW convention. The observed comparisons show that the adopted interpretation is not safe to use as a verified common-angle definition for this catalogue snapshot.')
table(['PA comparison','All 2207 fits','546 unflagged'],[
 ['Stored conversion median absolute difference', '82.93°','86.28°'],
 ['Raw catalogue angle treated as sky PA hypothesis','7.01°','3.78°'],
 ['Stored converted angle plus 90° diagnostic','7.07°','3.72°']], [4,1.4,1.4])
pic('pa_compact.png','Figure 2. Axial PA differences in [−90°, 90°). The quarter-turn panel is a diagnostic hypothesis, not an approved correction or a new catalogue measurement.')
p('For unflagged fits, only 2 of 546 are within 10° under the stored conversion, whereas 433 of 546 (79.3%) are within 10° under the quarter-turn diagnostic. The diagnostic 90th-percentile absolute difference is 15.25°. A systematic shift this large is stronger evidence of an angle-definition mismatch than of genuine disc directions being perpendicular to every catalogue detection ellipse.')
p('Two original diagnostic overlays and a flagged example were inspected. One unflagged galaxy, 2731221676651025355, has a fitted disc PA of 131.69° but a stored catalogue-converted PA of 41.69°, a 90.00° difference. Its native image and fitted ellipses support the fitted elongation. The evidence implicates the comparison convention; the image-based disc estimate still needs the usual physical review.')
h('How to use the PA columns now')
bullet('Retain the original raw angle and converted value for reproducibility; do not use their stored difference to rank physical outliers yet.')
bullet('Confirm the release-specific catalogue convention with the Euclid team, including the reference axis, rotation direction and whether it is already a sky angle.')
bullet('Test the confirmed convention against image ellipses across multiple tiles, then regenerate catalogue comparison columns without rerunning the expensive image fits.')
p('PA_convention_audit.csv preserves the original comparison and the two hypotheses. The completed source manifest was not modified. No blanket 90° correction has been silently applied, and the raw Sérsic PA is not substituted as a bar angle.')

page('Inclination comparison with Euclid shape proxies')
p('Both sets of inclination proxies use degrees with 0° face-on, assuming q0 = 0 and an intrinsically circular disc. The catalogue detection proxy uses q = 1 − ELLIPTICITY; the single-Sérsic proxy uses the fitted VIS b/a. These are whole-object shapes. Agreement is a consistency check, not an independent measurement of true disc inclination.')
table(['Comparison','All fits n 2207','Unflagged n 546'],[
 ['Detection proxy median absolute difference','4.28°','2.71°'],
 ['Detection proxy 90th-percentile absolute difference','15.37°','10.18°'],
 ['Detection proxy within 10°','1745 / 2207  79.1%','487 / 546  89.2%'],
 ['Sérsic proxy median absolute difference','8.41°','6.76°'],
 ['Sérsic proxy 90th-percentile absolute difference','23.69°','17.71°'],
 ['Sérsic proxy within 10°','1258 / 2207  57.0%','372 / 546  68.1%']], [3.7,1.55,1.55])
pic('inclination_compact.png','Figure 3. Calculated inclination versus Euclid whole-object proxies. Pale points are all provisional fits; dark points are the unflagged subset. The dashed diagonal marks equality.')
p('Signed medians, defined as calculated minus Euclid proxy, are +1.46° for detection and −5.04° for Sérsic across all fits. In the unflagged subset they are +0.26° and −3.77°. Thus the calculated disc tilt is typically closer to the detection proxy than the Sérsic proxy, and tends to be lower than the Sérsic-derived inclination.')
p('Different radial weighting can explain part of these differences: a single Sérsic fit blends bulge, bar and disc, while the algorithm selects an outer interval. The current analysis cannot separate this effect from masking, PSF, neighbour or range-selection biases. A finite intrinsic disc thickness would also change all inferred inclinations. Any new q0 should be applied consistently and recorded.')
p('No catalogue disc-fit b/a values are populated for these 2,777 sources, so there is no direct Euclid disc-inclination comparison here. The geometry should not be promoted to final science values solely because its detection proxy agrees.')

page('Zoobot bar classifications and Zenodo provenance')
p('The Euclid archive contains usable BAR_STRONG, BAR_WEAK and BAR_NO concentrations for 2,767 of the 2,777 galaxies. Ten have missing or invalid values and remain blank. Each valid concentration was divided by the sum of the three to obtain comparable strong, weak and no-bar fractions. Strong plus weak gives the any-bar fraction.')
p('Zenodo fractions are preserved separately from the archive values. The selection required Zenodo any-bar fraction ≥0.8, so it is expected that the Zenodo axis of that panel is truncated. These scores describe predicted classifications; they are not bar torque, force ratio, length, ellipticity or position angle.')
table(['Fraction comparison','Median signed difference','Median absolute difference'],[
 ['Strong bar','+0.0572','0.0853'],['Weak bar','−0.0599','0.0620'],['Any bar','+0.0037','0.0367']], [2,2.3,2.5])
pic('zoobot_compact.png','Figure 4. Euclid normalized fractions against Zenodo fractions for 2,767 matched objects. Differences are not unexpected identity failures; both the raw values and exact IDs are retained.')
p('The archive tends to assign a larger strong-bar fraction and smaller weak-bar fraction than Zenodo. The any-bar fraction is closer overall: its 90th-percentile absolute difference is 0.1124. Nevertheless, 535 of the 2,767 objects (19.3%) have archive any-bar fraction below 0.8 despite passing the Zenodo threshold. Applying that threshold to the archive would therefore define a different sample. The present audit does not establish which model or processing differences produced the disagreement.')
h('Version record')
p('The selection used Zenodo record 15106473, DOI 10.5281/zenodo.15106473, the fourth record in the retrieved version history. The metadata API does not supply a separate version label, so the manifest records ordinal 4 and leaves the version-label field empty. The record was updated on 29 March 2025. This report preserves that version-specific source rather than referring only to a moving latest version.')
p('The Zenodo published CSV checksum is md5:652e4546c63085c37539a9e9664f64e2. It is provenance metadata; the earlier streamed selection was not independently checked against a full downloaded-catalogue hash. The extracted FITS download batches were separately SHA-256 verified during local transfer.')

page('Illustrative image and profile checks')
p('These are selected diagnostic examples, not a statistical visual validation of all 2,207 fits. Cyan curves show fitted ellipses; orange and green shading mark the selected bar and disc radial intervals. Image panels use native pixel axes, whereas profiles show sky PA.')
pic('case_unflagged.png','Figure 5. Unflagged object 2731221676651025355. The inner and outer directions are fairly stable, but passing these tests does not independently prove the disc range is uncontaminated by bar light.')
p('Calculated bar PA = 141.85°, disc PA = 131.69° and inclination = 65.37°. The stored converted Euclid PA = 41.69° is nearly perpendicular. This is an example of the convention issue, not a defensible physical 90° disagreement. The bar and disc are also close in apparent direction, which makes separate component interpretation important.')
pic('case_flagged.png','Figure 6. Flagged object −515488559271594427. The outer PA varies strongly where the galaxy is close to circular; the automatic inner bar range is small.')
p('Calculated bar PA = 78.21°, disc PA = 151.15° and inclination = 27.25°. Flags identify unstable disc PA and sensitivity of the small bar range to resolution. The large outer PA changes show why a single returned angle should not be accepted merely because the fitter completed. This case needs masking and radial-range review.')
p('The previously inspected larger galaxy −547491640287288205 has an outer-disc PA of about 82.1° and inclination about 68.6°. Its raw catalogue angle is 81.1°, while the stored conversion is 171.2°. That additional image check is consistent with the same convention problem across a different tile and hemisphere.')

page('Recommended next steps and reproducibility')
h('Adopt geometry only after review')
bullet('Resolve the Euclid PA convention first and regenerate only the comparison fields. Do not rerun the ellipse fits merely to repair a catalogue conversion.')
bullet('Start physical validation with the 546 unflagged fits, checking the 19 partial boxes explicitly. Inspect the chosen disc and bar intervals, centres and connected neighbours.')
bullet('Recover the 431 radial-selection failures by reviewing profiles and supplying object-specific ranges. Review the remaining 139 failures for short profiles or the single centre/component mismatch.')
bullet('Treat the 1,210 small-bar flags as a resolution and model-selection concern. Consider unbinned refits and PSF-aware component modelling for candidates whose bar remains unresolved or ambiguous.')
bullet('Build an adopted-geometry table with explicit provenance, q0, radial ranges, flags and visual-review status before creating final deprojected images. Keep failed fits and missing catalogue values visible.')
p('All returned geometry remains provisional. This assessment has not created a scientifically approved deprojection catalogue or modified the native cutouts. A correct numerical transformation cannot compensate for a wrong physical disc or bar identification.')
h('Files supporting this report')
p('Primary data: MER_VIS_local_geometry_manifest.csv in the Euclid_MER_VIS_geometry data folder. Algorithm and run configuration: calculate_local_MER_VIS_geometry.py and MER_VIS_local_geometry_run.json. The geometry algorithm uses the portable Photutils fitter and MER_VIS_geometry_transforms.py. Diagnostics are in the data folder’s profiles and logs subfolders.')
p('The companion statistics.json records counts and comparison summaries. PA_convention_audit.csv retains a per-fit audit of competing angle interpretations. Original calculated angles, cached Euclid values, raw flags and Zenodo provenance remain in the source manifest. No authoritative catalogue values were overwritten for the report.')
p('Source manifest SHA-256')
a=doc.add_paragraph(s['source_sha256'])
for r in a.runs:r.font.name='Consolas';r.font.size=Pt(8)
h('References')
for text in [
'Euclid MER Morphology Cookbook — detection, Sérsic and Zoobot definitions. https://euclid.esac.esa.int/dr/q1/dpdd/merdpd/mermorphologycookbook.html',
'Euclid Final Catalog Product — column units and descriptions. https://euclid.esac.esa.int/dr/q1/dpdd/merdpd/dpcards/mer_finalcatalog.html',
'SExtractor shape conventions — image-axis angles and ellipticity. https://sextractor.readthedocs.io/en/latest/Position.html',
'Galaxy Zoo Euclid Q1 morphology catalogue — version-specific source. https://doi.org/10.5281/zenodo.15106473',
'Photutils elliptical isophote analysis — fitting method. https://photutils.readthedocs.io/en/latest/user_guide/isophote.html']:
 a=doc.add_paragraph(text)
 for r in a.runs:r.font.size=Pt(9)
doc.core_properties.title='Euclid MER VIS galaxy geometry fitting report';doc.core_properties.author=''
for element in (doc._element,doc.styles.element):
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(root/'Euclid_MER_VIS_geometry_fitting_report.docx')
