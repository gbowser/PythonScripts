#!/usr/bin/env python3
"""Build the SEP failure-mode analysis Word report."""
from __future__ import annotations

import csv
from collections import Counter
from datetime import datetime
import json
from pathlib import Path
from statistics import mean, median

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
CSV_PATH = ROOT / "Foreground Masking/Optimisation/sep_galaxy150_per_toy_recovery.csv"
EXP_ROOT = Path("/mnt/d/Dropbox/Public Documents/UCLAN/MSc Research/Remove foreground objects/clean22_haigh_aligned_bright150_galaxy150_sep_sensitivity")
SEP_CV = EXP_ROOT / "SEP_cross_validation/cross_validation_candidates.csv"
SEP_REJECTED = EXP_ROOT / "SEP_cross_validation/sep_toy_cross_validation_rejected.json"
MTO_CV = EXP_ROOT / "MTObjects_cross_validation/cross_validation_candidates.csv"
OUT = ROOT / "Foreground Masking/documentation/SEP Failure Mode Analysis - Haigh Aligned Galaxy150 Experiment.docx"
ASSETS = ROOT / "Foreground Masking/documentation/sep_failure_mode_assets"

BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
RED = "A61B1B"
GOLD = "9A6700"
GREEN = "207A46"
GRAY = "666666"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def pct(value: float, digits: int = 1) -> str:
    return f"{100 * value:.{digits}f}%"


def set_font(run, size=None, bold=None, color=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), "120")
    ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, data, widths, numeric_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(value))
        set_font(r, size=9.5, bold=True, color=DARK)
    for values in data:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=9.25)
    table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    p.add_run(text)


def add_callout(doc, label, text, color=DARK):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=color)
    r = p.add_run(text)
    set_font(r)
    table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=9, italic=True, color=GRAY)


def group_stats(rows, predicate):
    selected = [r for r in rows if predicate(r)]
    return len(selected), mean(float(r["detected"]) for r in selected), mean(float(r["recall"]) for r in selected)


def bin_stats(rows, field, bins):
    result = []
    for label, low, high in bins:
        selected = [r for r in rows if low <= float(r[field]) < high]
        result.append((label, len(selected), mean(float(r["detected"]) for r in selected), mean(float(r["recall"]) for r in selected)))
    return result


def figures(rows):
    ASSETS.mkdir(parents=True, exist_ok=True)
    stars = [r for r in rows if r["object_type"] == "star"]
    galaxies = [r for r in rows if r["object_type"] == "galaxy"]

    fig, ax = plt.subplots(figsize=(7.2, 3.7))
    x = np.arange(2)
    det = [mean(float(r["detected"]) for r in stars), mean(float(r["detected"]) for r in galaxies)]
    rec = [mean(float(r["recall"]) for r in stars), mean(float(r["recall"]) for r in galaxies)]
    ax.bar(x - .18, np.array(det) * 100, .36, label="Detection rate", color="#2E74B5")
    ax.bar(x + .18, np.array(rec) * 100, .36, label="Mean toy recall", color="#8DB3D3")
    ax.set_xticks(x, [f"IRAC-PSF stars\n(n={len(stars)})", f"Sérsic galaxies\n(n={len(galaxies)})"])
    ax.set_ylim(0, 60); ax.set_ylabel("Recovery (%)"); ax.legend(frameon=False, ncol=2, loc="upper right")
    ax.grid(axis="y", alpha=.2)
    fig.tight_layout(); fig.savefig(ASSETS / "type_recovery.png", dpi=180); plt.close(fig)

    modes = ["recovered", "not extracted", "filtered: area", "filtered: centre", "filtered: both", "baseline overlap"]
    def mode_counts(source_rows):
        c = Counter(r["failure_mode"] for r in source_rows)
        return [c["recovered"], c["not_extracted_or_undersegmented"], c["filtered:max_area"], c["filtered:protected_centre"], c["filtered:max_area+protected_centre"], c["baseline_overlap_removed_increment"]]
    fig, ax = plt.subplots(figsize=(7.2, 4.1))
    left = np.zeros(2)
    palette = ["#207A46", "#D9A441", "#D67447", "#B44C4C", "#7E2F2F", "#777777"]
    totals = np.array([len(stars), len(galaxies)])
    for label, color, vals in zip(modes, palette, np.array([mode_counts(stars), mode_counts(galaxies)]).T):
        percentages = 100 * vals / totals
        ax.barh([0, 1], percentages, left=left, label=label, color=color)
        left += percentages
    ax.set_yticks([0, 1], ["Stars", "Sérsic galaxies"]); ax.set_xlim(0, 100); ax.set_xlabel("Percentage of toys")
    ax.legend(frameon=False, fontsize=8, ncol=3, bbox_to_anchor=(.5, 1.24), loc="upper center")
    fig.tight_layout(); fig.savefig(ASSETS / "failure_modes.png", dpi=180, bbox_inches="tight"); plt.close(fig)

    radius = bin_stats(galaxies, "effective_radius_arcsec", [("<1.5", 0, 1.5), ("1.5–3", 1.5, 3), ("3–4.5", 3, 4.5), (">4.5", 4.5, 99)])
    sersic = bin_stats(galaxies, "sersic_index", [("2–2.5", 2, 2.5), ("2.5–3", 2.5, 3), ("3–3.5", 3, 3.5), ("3.5–4", 3.5, 4.01)])
    local = bin_stats(rows, "local_background_z", [("<−0.5", -99, -.5), ("−0.5–0", -.5, 0), ("0–0.5", 0, .5), (">0.5", .5, 99)])
    distance = bin_stats(rows, "distance_normalised", [("<0.5", 0, .5), ("0.5–0.75", .5, .75), ("0.75–1", .75, 1), (">1", 1, 99)])
    fig, axs = plt.subplots(2, 2, figsize=(7.2, 6.2))
    for ax, title, values in zip(axs.flat, ["Galaxy effective radius (arcsec)", "Galaxy Sérsic index", "Local background above outer-field median (σ)", "Normalised distance from target centre"], [radius, sersic, local, distance]):
        ax.bar([v[0] for v in values], [100 * v[2] for v in values], color="#2E74B5")
        ax.set_title(title, fontsize=10); ax.set_ylim(0, 75); ax.set_ylabel("Detection (%)"); ax.grid(axis="y", alpha=.2)
        ax.tick_params(axis="x", labelsize=8)
        for i, v in enumerate(values): ax.text(i, 100*v[2]+2, f"n={v[1]}", ha="center", fontsize=7)
    fig.tight_layout(); fig.savefig(ASSETS / "factor_detection.png", dpi=180); plt.close(fig)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    for level, size, color, before, after in ((1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,DARK,8,4)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"; style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(.5); style.paragraph_format.first_line_indent = Inches(-.25)
        style.paragraph_format.space_after = Pt(8); style.paragraph_format.line_spacing = 1.167


def footer_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("SEP failure-mode analysis  |  ")
    set_font(r, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def main():
    rows = read_csv(CSV_PATH)
    sep_candidates = read_csv(SEP_CV)
    rejection = json.loads(SEP_REJECTED.read_text(encoding="utf-8"))
    mto = read_csv(MTO_CV) if MTO_CV.exists() else []
    figures(rows)
    stars = [r for r in rows if r["object_type"] == "star"]
    galaxies = [r for r in rows if r["object_type"] == "galaxy"]
    outcomes = {kind: Counter(r["failure_mode"] for r in source) for kind, source in (("Stars", stars), ("Galaxies", galaxies))}

    doc = Document()
    configure_styles(doc)
    footer_page_number(doc.sections[0])
    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Foreground Masking Research | Technical Findings")
    set_font(r, size=9, color=GRAY)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    r = p.add_run("TECHNICAL FINDINGS"); set_font(r, size=10, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SEP Failure-Mode Analysis"); set_font(r, size=25, bold=True, color=DARK)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Haigh-aligned foreground-object experiment: 1.5× Sérsic-galaxy scale, unchanged IRAC-PSF stars"); set_font(r, size=12.5, color=GRAY)
    add_table(doc, ["Analysis basis", "Value"], [
        ["Science sample", "22 visually clean S4G galaxy frames"],
        ["Saved injections", "3 training seeds + 2 validation seeds"],
        ["Toy population", "375 sources: 263 stars and 112 Sérsic galaxies"],
        ["SEP candidate", "Best diagnostic candidate from held-out fold 20"],
        ["Report generated", datetime.now().strftime("%-d %B %Y, %H:%M")],
    ], [2400, 6960])
    add_callout(doc, "Principal finding", "SEP is missing both source classes, with Sérsic galaxies performing worst. The dominant failure is component merging with target-galaxy emission, followed by rejection under the maximum-area and protected-centre filters—not inadequate toy brightness.", RED)

    add_heading(doc, "1. Question and analytical approach")
    doc.add_paragraph("The analysis asked whether SEP’s failed optimisation was caused primarily by stars, Sérsic galaxies, brightness, morphology, placement, or post-extraction filtering. The audit reapplied the best diagnostic parameter set to every saved injection in all five seed sets. Each toy was followed through four stages: raw SEP segmentation, component filtering, final dilation, and subtraction of the baseline mask.")
    add_bullet(doc, "Toy detection means that at least 50% of that toy’s truth pixels were newly masked relative to the uninjected baseline.")
    add_bullet(doc, "Local background was measured from the original uninjected frame and expressed relative to the outer-field robust noise estimate.")
    add_bullet(doc, "Distance was normalised by the equivalent radius of the displayed analysis frame; it is a placement diagnostic, not a physical galactocentric radius.")

    doc.add_page_break()
    add_heading(doc, "2. Recovery by source class")
    type_rows = []
    for label, source in (("IRAC-PSF stars", stars), ("PSF-convolved Sérsic galaxies", galaxies)):
        type_rows.append([label, len(source), pct(mean(float(r["detected"]) for r in source)), pct(mean(float(r["recall"]) for r in source)), pct(median(float(r["recall"]) for r in source))])
    add_table(doc, ["Source class", "n", "Detection", "Mean recall", "Median recall"], type_rows, [3300, 900, 1650, 1650, 1860], numeric_cols=(1,2,3,4))
    doc.add_picture(str(ASSETS / "type_recovery.png"), width=Inches(6.35))
    add_caption(doc, "Figure 1. Per-toy recovery for the best diagnostic SEP parameter set across all five saved seed sets.")
    doc.add_paragraph("Stars are recovered about 12.6 percentage points more often than Sérsic galaxies, but neither class reaches the required 50% detection rate. The zero median for both classes indicates a largely bimodal process: many toys are substantially recovered, while at least half receive no incremental masking at all.")

    add_heading(doc, "3. Where the toys are lost")
    failure_rows = []
    ordered = [
        ("Recovered", "recovered"),
        ("Not extracted / raw overlap <50%", "not_extracted_or_undersegmented"),
        ("Rejected: maximum area", "filtered:max_area"),
        ("Rejected: protected-centre centroid", "filtered:protected_centre"),
        ("Rejected: both area and centre", "filtered:max_area+protected_centre"),
        ("Removed by baseline-overlap rule", "baseline_overlap_removed_increment"),
    ]
    for label, key in ordered:
        failure_rows.append([label, outcomes["Stars"][key], pct(outcomes["Stars"][key]/len(stars)), outcomes["Galaxies"][key], pct(outcomes["Galaxies"][key]/len(galaxies))])
    add_table(doc, ["Outcome", "Stars n", "Stars %", "Galaxies n", "Galaxies %"], failure_rows, [3900, 1200, 1320, 1500, 1440], numeric_cols=(1,2,3,4))
    doc.add_picture(str(ASSETS / "failure_modes.png"), width=Inches(6.35))
    add_caption(doc, "Figure 2. Classification of each toy at the point where recovery succeeded or failed.")
    add_callout(doc, "Mechanism", "For 132 of 263 stars and 58 of 112 galaxies, SEP initially extracted a component at the toy position but later rejected it because the component was too large and/or its centroid lay inside the protected centre. Since toys were not placed in the nucleus, a protected-centre rejection generally means the toy segment merged with target-galaxy structure and inherited a centroid closer to the galaxy centre.", GOLD)
    doc.add_paragraph("Only 18 of 375 toys failed primarily because SEP did not extract or sufficiently segment them. This makes a simple increase in brightness or detection aggressiveness unlikely to solve the principal problem. Increasing aggressiveness can instead enlarge and connect components, worsening the area/centre rejection and masking burden.")

    add_heading(doc, "4. Brightness and morphology")
    brightness_rows = []
    bins = [("9–18σ", 9, 18), ("18–30σ", 18, 30), ("30–46σ", 30, 46)]
    for label, low, high in bins:
        sn, sd, sr = group_stats(stars, lambda r, lo=low, hi=high: lo <= float(r["peak_sigma"]) < hi)
        gn, gd, gr = group_stats(galaxies, lambda r, lo=low, hi=high: lo <= float(r["peak_sigma"]) < hi)
        brightness_rows.append([label, sn, pct(sd), pct(sr), gn, pct(gd), pct(gr)])
    add_table(doc, ["Peak", "Stars n", "Stars det.", "Stars recall", "Gal. n", "Gal. det.", "Gal. recall"], brightness_rows, [1260, 1080, 1320, 1440, 1080, 1500, 1680], numeric_cols=(1,2,3,4,5,6))
    doc.add_paragraph("Star recovery improves from the faintest bin and then plateaus. Galaxy recovery declines in the brightest bin. This counter-intuitive galaxy result is consistent with brighter extended wings making a source more likely to connect to surrounding galaxy emission, rather than with an insufficient signal-to-noise ratio.")
    radius = bin_stats(galaxies, "effective_radius_arcsec", [("<1.5", 0, 1.5), ("1.5–3", 1.5, 3), ("3–4.5", 3, 4.5), (">4.5", 4.5, 99)])
    sersic = bin_stats(galaxies, "sersic_index", [("2–2.5", 2, 2.5), ("2.5–3", 2.5, 3), ("3–3.5", 3, 3.5), ("3.5–4", 3.5, 4.01)])
    morph_rows = [["Effective radius", *v] for v in radius] + [["Sérsic index", *v] for v in sersic]
    add_table(doc, ["Factor", "Bin", "n", "Detection", "Mean recall"], [[a,b,c,pct(d),pct(e)] for a,b,c,d,e in morph_rows], [2200, 1800, 900, 2100, 2360], numeric_cols=(2,3,4))
    doc.add_paragraph("The largest galaxies (>4.5 arcsec effective radius after scaling) have only 22.2% detection. High-concentration sources with Sérsic index 3.5–4 perform worst at 11.5%. Axis ratio shows no comparable monotonic effect, so elongation is not the primary morphology driver in this experiment.")

    add_heading(doc, "5. Placement and local environment")
    doc.add_picture(str(ASSETS / "factor_detection.png"), width=Inches(5.45))
    add_caption(doc, "Figure 3. Detection rates by Sérsic morphology and by two environment diagnostics. Bin counts are shown above each bar.")
    doc.add_paragraph("Recovery is strongly location-dependent even though all toy centres were restricted to the designated quiet-placement region. Detection rises from 21.2% inside a normalised distance of 0.5 to 58.7% beyond 1.0. It falls from 67.8% in the lowest local-background bin to 33.0% where the original local background is more than 0.5σ above the outer-field median.")
    add_callout(doc, "Interpretation", "The quiet-placement rule removes obvious target structure and pre-existing compact sources, but faint connected galaxy emission remains. SEP’s segmentation can bridge from the injected source through that low-level emission toward the target. This explains why a toy can be physically outside the protected nucleus yet be rejected as a centre-associated component.")

    doc.add_page_break()
    add_heading(doc, "6. Cross-validation feasibility")
    all_det = [float(r["all22_toy_detection_rate"]) for r in sep_candidates]
    all_rec = [float(r["all22_mean_toy_recall"]) for r in sep_candidates]
    mean_mask = [float(r["all22_mean_masked_fraction"]) for r in sep_candidates]
    max_mask = [float(r["all22_max_masked_fraction"]) for r in sep_candidates]
    add_table(doc, ["Across 22 fold candidates", "Minimum", "Maximum", "Required"], [
        ["All-22 toy detection", pct(min(all_det)), pct(max(all_det)), "≥50%"],
        ["All-22 mean toy recall", pct(min(all_rec)), pct(max(all_rec)), "≥30%"],
        ["All-22 mean displayed-frame mask", pct(min(mean_mask)), pct(max(mean_mask)), "Minimise"],
        ["Worst displayed-frame mask", pct(min(max_mask)), pct(max(max_mask)), "15% control"],
    ], [4200, 1680, 1680, 1800], numeric_cols=(1,2,3))
    doc.add_paragraph(f"No candidate achieved 50% toy detection. The best diagnostic candidate reached {pct(float(rejection['best_diagnostic_candidate']['all22_toy_detection_rate']))} detection and {pct(float(rejection['best_diagnostic_candidate']['all22_mean_toy_recall']))} mean toy recall, but masked {pct(float(rejection['best_diagnostic_candidate']['all22_mean_masked_fraction']))} of the displayed frame on average and {pct(float(rejection['best_diagnostic_candidate']['all22_max_masked_fraction']))} in the worst case. The rejection is therefore scientifically appropriate.")

    add_heading(doc, "7. Conclusions and recommended experiments")
    add_bullet(doc, "Treat component merging—not toy faintness—as the primary SEP failure mode. Further brightness increases are not recommended.")
    add_bullet(doc, "Test detection on a Gaussian residual or another target-suppressed image while applying the resulting mask to the original image. This directly addresses bridges through smooth target-galaxy light.")
    add_bullet(doc, "Audit deblending before relaxing maximum-area or protected-centre rules. Relaxing those filters alone would admit the very large galaxy-connected components responsible for 17–83% frame masking.")
    add_bullet(doc, "Retain quiet-region placement for the controlled optimisation, then add a separate stratified challenge set covering quiet background, faint outer structure and moderate target structure. Do not distribute toys uniformly over the entire frame or protected nucleus.")
    add_bullet(doc, "If additional statistical power is needed, add reproducible seeds rather than increasing the number of toys per frame. Use a balanced diagnostic seed set while preserving the scientifically motivated star/galaxy mixture in the primary experiment.")
    add_bullet(doc, "Report and gate stars and galaxies separately so stronger star recovery cannot conceal poor galaxy recovery.")

    doc.add_page_break()
    add_heading(doc, "8. Interim MTObjects indication (provisional)")
    if mto:
        md = [float(r["all22_toy_detection_rate"]) for r in mto]
        mr = [float(r["all22_mean_toy_recall"]) for r in mto]
        mm = [float(r["all22_mean_masked_fraction"]) for r in mto]
        mx = [float(r["all22_max_masked_fraction"]) for r in mto]
        hd = [float(r["held_out_toy_detection_rate"]) for r in mto]
        feasible = sum(int(float(r["candidate_feasible"])) for r in mto)
        add_table(doc, ["Metric", "Interim result"], [
            ["Completed folds at snapshot", f"{len(mto)} / 22"],
            ["Candidates marked feasible", f"{feasible} / {len(mto)}"],
            ["All-22 toy detection", f"mean {pct(mean(md))}; range {pct(min(md))}–{pct(max(md))}"],
            ["All-22 mean toy recall", f"mean {pct(mean(mr))}; range {pct(min(mr))}–{pct(max(mr))}"],
            ["Mean displayed-frame mask", f"mean {pct(mean(mm))}; range {pct(min(mm))}–{pct(max(mm))}"],
            ["Worst-case mask by candidate", f"range {pct(min(mx))}–{pct(max(mx))}"],
            ["Held-out toy detection", f"mean {pct(mean(hd))}; range {pct(min(hd))}–{pct(max(hd))}"],
        ], [3600, 5760])
        add_callout(doc, "Interim assessment", "MTObjects is currently materially stronger than SEP: its completed-fold candidates exceed the 50% all-22 detection and 30% recall gates while maintaining much lower mean masking. Held-out performance is variable, so the conclusion remains provisional until all 22 folds and final winner selection are complete.", GREEN)
    else:
        doc.add_paragraph("No MTObjects cross-validation candidates were available when the report was generated.")

    doc.add_page_break()
    add_heading(doc, "Appendix A. Reproducibility and definitions")
    add_table(doc, ["Item", "Definition"], [
        ["Experiment", "bright150_galaxy150: 9–45σ peaks; star PSF unchanged; Sérsic-galaxy scale increased by 50%"],
        ["Analysis population", "All 375 materialised toys from training seeds 1–3 and validation seeds 1–2"],
        ["Best diagnostic SEP parameters", rejection["best_diagnostic_candidate"]["parameter_set_json"]],
        ["Raw recall", "Fraction of toy truth pixels belonging to any raw SEP segment"],
        ["Filtered recall", "Fraction remaining after maximum-area, elongation and protected-centre component filtering"],
        ["Incremental recall", "Fraction newly masked after removing pixels already masked in the uninjected baseline"],
        ["Local-background z", "9×9-pixel local median minus outer-field median, divided by outer-field robust σ"],
    ], [2640, 6720])
    doc.add_paragraph("The detailed per-toy audit is retained as sep_galaxy150_per_toy_recovery.csv, and the reproducible audit code is analyse_sep_toy_type_recovery.py in the Optimisation folder.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SEP Failure-Mode Analysis – Haigh-Aligned Galaxy150 Experiment"
    doc.core_properties.subject = "Foreground-object masking optimisation"
    doc.core_properties.author = "MSc Research – Foreground Masking"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
