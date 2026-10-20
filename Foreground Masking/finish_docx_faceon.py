from pathlib import Path
from docx import Document
from docx.shared import Pt
p=Path('MER_VIS_summary_build/output/Euclid_MER_VIS_outcomes_summary.docx');d=Document(p)
d.styles['Normal'].font.size=Pt(10)
for a in d.paragraphs:
 if a.text=='Fitting outcomes':a.text='Fitting outcomes under the original screening policy'
for t in d.tables:
 if t.rows[0].cells[0].text=='Quality condition':t.rows[0].cells[0].text='Original condition or informational statistic'
d.save(p)
