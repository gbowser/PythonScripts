from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
src=Path('update_summary_PA.py').read_text(encoding='utf-8-sig')
addition=src.split('addition="""',1)[1].split('"""',1)[0]
path=Path('MER_VIS_summary_build/output/Euclid_MER_VIS_outcomes_summary.docx')
doc=Document(path)
anchor=next(p for p in doc.paragraphs if p.text=='Comparison with Euclid geometry')
def insert(t,style=None):
 el=OxmlElement('w:p'); anchor._p.addprevious(el); p=Paragraph(el,anchor._parent);p.text=t
 if style:p.style=style
 return p
def page(t):insert(t,'Heading 1')
def h(t):insert(t,'Heading 2')
def p(t):insert(t)
exec(addition)
# Existing page break now starts the new methodology page; add a new break before comparison.
insert('').add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
for p in doc.paragraphs:
 if '19 October 2026' in p.text:p.text=p.text.replace('19 October 2026','Updated 20 October 2026')
doc.save(path)
# Perform only the PPTX builder edits from the update script.
rest=src.split("p=Path('MER_VIS_summary_build/build_summary.mjs');",1)[1]
rest=rest.replace('p.read_text()','p.read_text(encoding="utf-8")').replace('p.write_text(s)','p.write_text(s,encoding="utf-8")')
p=Path('MER_VIS_summary_build/build_summary.mjs')
exec(rest)
