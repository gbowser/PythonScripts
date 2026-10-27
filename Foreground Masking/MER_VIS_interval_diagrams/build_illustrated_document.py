from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_BREAK
SRC=Path(r'C:\Users\gordo\Documents\Github\PythonScripts\Euclid\reports\MER_VIS_current\Euclid_MER_VIS_disc_PA_and_inclination_explained_clarified.docx')
OUT=Path('MER_VIS_interval_diagrams');d=Document(SRC);ps=d.paragraphs;m=json.loads((OUT/'measured_example.json').read_text())
# Replace only the former page-three selection section. Retain inclination and PA prose below.
start=next(p for p in ps if p.text=='Disc and bar selection and the angle conventions')
end=next(p for p in ps if p.text=='How to interpret warnings')
incl=next(p.text for p in ps if p.text.startswith('For an intrinsically circular thin disc'))
pa=next(p.text for p in ps if p.text.startswith('Photutils pixel angles start'))
corr=next(p.text for p in ps if p.text.startswith('The additional +90'))
node=start._p
while node is not end._p:
 nxt=node.getnext();node.getparent().remove(node);node=nxt
anchor=end._p
# New blocks are created at the end and moved before the retained next section.
def p(text='',style=None,boldlead=None):
 q=d.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):q.add_run(boldlead).bold=True;q.add_run(text[len(boldlead):])
 else:q.add_run(text)
 anchor.addprevious(q._p);return q
def heading(text,level=1):return p(text,'Heading '+str(level))
def pagebreak():p().add_run().add_break(WD_BREAK.PAGE)
def figure(name):
 q=p();q.paragraph_format.space_before=Pt(3);q.paragraph_format.space_after=Pt(5);q.add_run().add_picture(str(OUT/name),width=Inches(6.45));return q
heading('How the disc interval and bar plateau are chosen')
p('An isophote fit describes one ellipse at one radius. Each ellipse contributes one ellipticity point and one PA point at that radius on the two profiles. The colours link the image to the graphs: orange is the candidate bar interval; green is the selected outer-disc interval. The diagram below uses invented values to make the definitions clear.')
figure('disc_bar_concept.png')
p('r_out is the semi-major axis of the outermost accepted ellipse in the selected bright profile. It is the outer bound available for disc selection, not the edge of the FITS cutout or a guaranteed physical boundary of the galaxy. r_bar is the semi-major axis of the largest ellipse in the accepted bar plateau. Both are distances from an ellipse centre to one end of its long axis, in arcseconds; neither is a diameter.',boldlead='r_out')
p('The disc interval starts at the larger of half r_out and 1.25 times r_bar, and ends at r_out. In this illustration, r_out = 10″ and r_bar = 3″. Half the outer radius is 5″; 1.25 times the bar end is 3.75″. We therefore use accepted disc ellipses from 5″ to 10″. The gap after the bar helps avoid measuring the bar as part of the disc.',boldlead='The disc interval')
p('The bar plateau is the orange cluster of high-ellipticity ellipses near a local peak. Their orientations must also be similar, as shown in the lower graph. A plateau need not be perfectly flat: small changes in ellipticity and PA are allowed. Moving outwards, the ellipses become rounder, so ellipticity falls. Here the drop is 0.60 − 0.35 = 0.25. High ellipticity alone is insufficient evidence for a bar.',boldlead='The bar plateau')
pagebreak();heading('A measured example of the same definitions')
p('Galaxy −639567775462084647 has five selected bar ellipses and six selected disc ellipses, with no counted geometry flags in the stored run. The image is the native observed cutout, not deprojected or rotated. The annotations below were calculated from its saved profile; no fit was rerun for this explanation.')
figure('disc_bar_measured.png')
p('Orange circles on the graphs correspond to the five orange ellipses on the image. Green circles and ellipses identify the selected disc. Blue denotes other accepted measurements; grey denotes excluded measurements. The purple dashed curve is the three-point running median used to locate an ellipticity peak among the accepted search samples. It is not a separate ellipse fit.')
table=d.add_table(rows=1, cols=3);anchor.addprevious(table._tbl)
for c,text in zip(table.rows[0].cells,['Bar ellipse radius (arcsec)','Ellipticity','Sky PA (degrees)']):c.text=text
for v in m['bar_points']:
 cells=table.add_row().cells
 for c,text in zip(cells,[f"{v['sma_arcsec']:.3f}",f"{v['ellipticity']:.3f}",f"{v['pa_sky_deg']:.2f}"]):c.text=text
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
table._tbl.tblPr.append(borders)
for r in table.rows:
 for c in r.cells:
  for q in c.paragraphs:
   q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2)
for c in table.rows[0].cells:
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF3');c._tc.get_or_add_tcPr().append(shade)
 for run in c.paragraphs[0].runs:run.bold=True
p('The bar PAs vary from 43.62° to 48.11°, with axial RMS scatter of 1.55°. This meets the 10° common bar threshold and the stronger 5° class A threshold. The adopted bar PA is the error-weighted axial mean, 46.34°. These neighbouring ellipses are correlated; five measurements are not five independent confirmations of a bar.')
p('The smoothed peak is 0.662, and the lowest accepted ellipticity in the outward comparison region is 0.402 at radius 2.422″. The recorded drop is therefore 0.260. The highest raw orange point is 0.672; it is not the smoothed peak used in that subtraction.')
p('Here r_out = 3.812″ and r_bar = 1.724″. The disc starts at max(1.906″, 2.155″) = 2.155″. Because fits exist only at discrete radii, its first selected ellipse is at 2.163″; six accepted disc ellipses extend to 3.812″. Their adopted disc PA is 38.60°. The bar and disc directions can differ.')
pagebreak();heading('The bar acceptance test and the reported angles')
heading('How the plateau and outward drop are tested',2)
p('The search examines accepted ellipses between 0.5 arcsec and 0.7 times r_out. It smooths their ellipticity values with a three-point median and looks for local peaks of at least 0.20. Around a peak, it extends the plateau while the smoothed ellipticity stays within 0.05 below the peak, with a limit of 12 neighbouring samples on either side. At least three accepted ellipses and plateau PA RMS scatter no greater than 10° are required.')
p('The outward drop is the smoothed peak ellipticity minus the lowest accepted ellipticity in a bounded region beyond that peak. The region extends to the larger of 2.5 times the peak radius or the peak radius plus 1 arcsec. The peak radius is where the candidate peak occurs; it need not equal r_bar, which is the plateau’s outer end. The code uses accepted measurements in this region, not an arbitrarily chosen faint outer pixel or a grey excluded point.')
p('The usual requirement is a drop of at least 0.08. A smaller drop of at least 0.04 can pass if either the outer PA twist or the difference from the preliminary outer-disc PA is at least 10°. Class A additionally requires peak ellipticity at least 0.30, a drop of at least 0.10 and plateau PA scatter no greater than 5°. Accepted candidates missing one or more stronger conditions are class B. The measured example passes all three class A conditions.')
p('A uniformly elongated disc can have high ellipticity and a stable PA without a separate bar plateau and outward drop. Conversely, a peak with strongly changing PA can be caused by spiral structure rather than a bar. Even a passing class A signature remains a candidate: a bulge, ring, arm or contaminant can imitate this profile. These numerical thresholds are operational choices in our pipeline, not universal physical definitions of a bar.')
heading('Inclination',2);p(incl)
heading('PA convention and mosaic orientation',2);p(pa);p(corr)
pagebreak()
# Correct a field-name ambiguity exposed while checking the radius examples.
for q in d.paragraphs:
 if q.text.startswith('r_out is the maximum accepted bright-profile semi-major axis, stored as codex_radius_arcsec.'):
  q.text=q.text.replace('r_out is the maximum accepted bright-profile semi-major axis, stored as codex_radius_arcsec.', 'r_out is the maximum accepted bright-profile semi-major axis used for selection. In a normal, non-fallback disc selection, it equals codex_disc_max_arcsec. It must not be confused with the separate view or extraction radius: the master run’s codex_radius_arcsec field serves that purpose.')
d.core_properties.comments='Illustrated disc/bar interval definitions and verified measured example added 26 October 2026. Original fit results unchanged.'
output=OUT/'Euclid_MER_VIS_disc_PA_and_inclination_explained_illustrated.docx';d.save(output)
print(output.resolve());print('Images:',len(d.inline_shapes))
