from pathlib import Path
from docx import Document
p=Path('MER_VIS_summary_build/output/Euclid_MER_VIS_outcomes_summary.docx');d=Document(p)
for t in d.tables:
 for row in t.rows:
  if row.cells[0].text=='Nearly face on  q > 0.90':
   row.cells[0].text='Nearly face on  q > 0.90 (informational statistic)'
for a in d.paragraphs:
 if a.text.startswith('Most detected issues arise'):
  a.text='Most detected issues arise from a small selected bar interval, unstable disc orientation or changing outer ellipticity. The table preserves the original run statistics. Nearly face-on is an informational statistic, not a quality issue in itself. Other flags identify specific numerical conditions and may overlap; their counts must not be added to infer a number of affected galaxies.'
 if a.text.startswith('No high-inclination or axis-ratio'):
  a.text='No high-inclination or axis-ratio-below-thickness flags occurred. A small-bar warning indicates sensitivity to resolution and interval selection; its interval boundary is not a measured physical bar length. Disc PA instability warns that the selected outer contours do not share a consistent direction, except that this is expected for nearly circular discs. Changing disc ellipticity makes inclination sensitive to the selected range.'
 if a.text.startswith('The unflagged fits have a median inclination'):
  a.text='Nearly face-on discs (200 of 2207 fits, 9.1%) have b/a > 0.90, corresponding to inclination below about 25.8° for a thin circular disc. Their disc PA is poorly constrained because the projected shape is almost circular. This does not make the galaxy or bar measurement poor: little deprojection is needed, and bar analysis can be easier. We therefore exclude this informational marker from revised quality counts and disregard disc PA instability for these 200 galaxies. Of the original 791 unstable disc PAs, 182 are face-on; the revised statistic is 609 of 2007 non-face-on fits (30.3%).'
 if a.text.startswith('QI includes any geometry flag'):
  a.text='The existing QI/NoQI report folders preserve the original screening policy: QI includes any original geometry flag, a failed fit or partial coverage. Its 2250 reports comprise 1661 originally flagged fits, 570 failures and 19 otherwise unflagged partial cases. These stored labels have not been regenerated. Under the revised interpretation used in the PowerPoint, nearly face-on is informational and unstable disc PA is disregarded for face-on galaxies. This gives 601 fits with no counted quality flags and 1606 with counted flags: 963 with one, 466 with two, 176 with three and one with four. Failures remain 570. These revised chart groups differ from the original QI/NoQI file groups.'
 if a.text.startswith('The 546-object comparison subset'):
  a.text='The comparison table retains the original 546-object subset with no original geometry flags, including 19 partial cases. It differs from both the 527 original NoQI reports and the revised 601 fits with no counted quality flags; the comparison statistics have not been recalculated for that revised subset. The completed validation of 51 stable non-diagonal discs supports the adopted PA adjustment, with a median absolute difference of 2.15°.'
 if a.text.startswith('Begin with the QI PDFs'):
  a.text='Begin with the QI PDFs and use the image and profiles together. Check the centre, neighbouring light and the selected outer interval. Nearly circular contours can have an ill-defined disc PA even when the fit completes. Treat the face-on marker as information rather than a quality defect, and disregard disc PA instability for face-on objects when assessing the revised quality counts. A face-on galaxy may remain useful for bar analysis.'
d.save(p)
