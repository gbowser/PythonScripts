from pathlib import Path
from docx import Document
root=Path(r'C:\Users\gordo\Documents\Github\PythonScripts\Euclid\reports\MER_VIS_current')
out=Path(r'C:\Users\gordo\Documents\Github\PythonScripts\Foreground Masking\MER_VIS_pipeline_v2\bar_class_update');out.mkdir(exist_ok=True)
sections=[('What class A and class B mean','These are operational categories defined by our ellipse-fitting code. They describe how clearly the radial ellipse profiles resemble a bar. They are not Galaxy Zoo or Zoobot classes, probabilities of a bar, or measurements of physical bar strength.'),('The common acceptance test','A candidate must have a smoothed ellipticity peak of at least 0.20 and a plateau of at least three accepted ellipses. The plateau sky PA scatter must be no more than 10°. Beyond the peak, ellipticity must drop by at least 0.08; alternatively, a drop of at least 0.04 is accepted if the outer PA twist or difference from the preliminary outer-disc PA is at least 10°.'),('Class A means a clearer ellipse signature','An accepted candidate is class A only if all three stronger conditions hold: ellipticity drops by at least 0.10 beyond the peak; plateau PA scatter is no more than 5°; and peak ellipticity is at least 0.30. Ellipticity is 1 − b/a, so 0.30 corresponds to an ellipse axis ratio b/a of 0.70.'),('Class B means accepted with weaker evidence','A candidate is class B if it passes the common acceptance test but misses at least one class A threshold. For example, peak ellipticity 0.35, drop 0.08 and PA scatter 4° passes the common test but is B because the drop is below 0.10. Peak ellipticity 0.40, drop 0.15 and scatter 7° is also B because scatter exceeds 5°.'),('How to interpret the quality flag','The flag weak_ellipse_bar_signature marks every accepted class B candidate for review. Class B does not mean that the galaxy definitely has a physically weak bar. A bulge, ring or spiral arm can mimic a bar profile, including a class A profile. Inspect the image and selected ellipses before adopting the bar PA. If no candidate passes the common test, no bar PA is adopted; the disc measurement may still be retained.')]
for name in ['Euclid_MER_VIS_outcomes_summary.docx','Euclid_MER_VIS_disc_PA_and_inclination_explained.docx']:
 d=Document(root/name)
 for p in d.paragraphs:
  if 'The bar passes the acceptance tests but is class B rather than A.' in p.text:
   p.text=p.text.replace('The bar passes the acceptance tests but is class B rather than A.','An accepted class B candidate meets the broader bar-profile test but misses one or more class A thresholds. See the class A and B explanation at the end of this document.')
 d.add_page_break()
 d.add_heading('Understanding the bar signature classes',level=1)
 for title,body in sections:
  d.add_heading(title,level=2);d.add_paragraph(body)
 d.save(out/name)
