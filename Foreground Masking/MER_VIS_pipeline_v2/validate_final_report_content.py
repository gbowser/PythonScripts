from pathlib import Path
import json,csv,hashlib,zipfile
from docx import Document
from pptx import Presentation
root=Path('MER_VIS_pipeline_v2/current_report/output');docfiles=['Euclid_MER_VIS_outcomes_summary.docx','Euclid_MER_VIS_disc_PA_and_inclination_explained.docx']
for name in docfiles:
 d=Document(root/name);text='\n'.join(p.text for p in d.paragraphs)
 assert 'previous 570' not in text and 'original run had' not in text.lower() and 'recovered interval failures' not in text.lower()
 assert 'six' in text and '227' in text and 'Class A' in text
 print(name,'checked')
ppt=root/'Euclid_MER_VIS_outcomes_summary_reviewed.pptx';p=Presentation(ppt);text='\n'.join(s.text for sl in p.slides for s in sl.shapes if s.has_text_frame)
assert len(p.slides)==14 and 'Recovered interval failures' not in text and 'Recovered short-profile failures' not in text and '570 original' not in text
print('PowerPoint',len(p.slides),'slides, historical comparison slides absent')
