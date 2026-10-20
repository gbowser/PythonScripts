from docx import Document
p=r'D:\Dropbox\Public Documents\UCLAN\MSc Research\Euclid\Claude Method\Euclid VIS disc PA, bar PA and inclination methodology.docx'
d=Document(p)
for a in d.paragraphs:
 if a.text.strip():print(a.text)
for t in d.tables:
 print('TABLE')
 for r in t.rows:print(' | '.join(c.text for c in r.cells))
