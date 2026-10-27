from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_BREAK
src=Path(r'C:\Users\gordo\Documents\Github\PythonScripts\Euclid\reports\MER_VIS_current\Euclid_MER_VIS_disc_PA_and_inclination_explained.docx')
out=Path('MER_VIS_acceptance_doc_update')/src.name
d=Document(src);ps=d.paragraphs
p=ps[6]
old='Here area is the catalogue segmentation footprint in native image pixels squared, pi is the circle constant (approximately 3.14159), and sqrt means square root.'
new='Here area is the number of native image pixels assigned to the detected source above the catalogue analysis threshold. It is read from the CSV manifest (sourced from the MER source catalogue, via the Zenodo catalogue), not from the science-image FITS header. This is the detected segmentation footprint, not the cutout rectangle or necessarily all the faint outer light. The quantity represents area in native pixel units; pi is the circle constant (approximately 3.14159), and sqrt means square root.'
assert old in p.text;p.text=p.text.replace(old,new)
# Give the expanded explanation its own page so the existing examples stay intact.
break_p=ps[8].insert_paragraph_before();break_p.add_run().add_break(WD_BREAK.PAGE)
ps[8].text='Which fitted isophotes are retained';ps[8].style='Heading 1'
ps[9].text='Acceptance here refers to one fitted isophote: the ellipse measured at one semi-major-axis radius. That ellipse has its own centre, ellipticity, PA, brightness and numerical diagnostics. The code applies all of the strict tests below to each ellipse. It does not apply these tests just to the galaxy centre, and an accepted ellipse is not automatically identified as the physical disc or bar.'
ps[10].text=''
items=[
('Normal numerical convergence. ', 'Photutils must report a numerically valid result and normal convergence (recorded internally as valid=True and stop_code=0). This describes the numerical fitting outcome; it does not establish that the ellipse traces the intended physical structure.'),
('Usable sampling of at least 65%. ', 'At least 65% of the sampled positions along that ellipse must supply usable intensity measurements after exclusions. For example, 65 usable positions out of 100 meets this test; 64 does not. This is not the percentage of the galaxy area covered, and it is separate from how many radial ellipses are available.'),
('Estimated ellipticity error below 0.1. ', 'Ellipticity is 1 − b/a. The 0.1 limit is an absolute, dimensionless uncertainty reported by the fit, not a PA error and not a 10% relative error. An ellipticity of 0.40 with a reported error of 0.08 passes this particular test; an error of 0.12 does not.'),
('Estimated PA error below 10 degrees. ', 'The reported fitting uncertainty in this ellipse’s orientation must be less than 10°. The code converts Photutils’ PA error from radians to degrees. This is a per-ellipse error estimate, distinct from PA RMS scatter across several selected disc ellipses.'),
('Positive brightness and limited centre movement. ', 'The background-subtracted fitted intensity must be positive. The ellipse centre must stay within the larger of three fitting-image pixels or 15% of its semi-major axis from the current reference centre. These distances use the same fitting-image pixel units. The reference can be updated during recovery. This test limits displacement; it does not independently verify the nucleus.'),
('Where the thresholds came from. ', 'The 65%, 0.1 and 10° limits, and the centre-displacement limits, were chosen in our pipeline code as practical screening rules. They are not Euclid requirements, Photutils defaults or published universal acceptance standards. No completed calibration study establishes these exact cutoffs as optimal for this sample. They need validation using representative visually reviewed cases and tests of sensitivity to alternative thresholds.'),
('Recovery and final galaxy measurements. ', 'The code normally seeks at least eight strict accepted ellipses in the overall profile. Labelled relaxed recovery permits 50% usable sampling, ellipticity error below 0.2, PA error below 20° and a larger centre tolerance. It can include fits that reached the iteration limit without normal convergence (stop code 2); selected examples raise iteration_limited_samples. Diverged or gradient-failure results (codes 4/5) are not treated as converged geometry. At least three accepted ellipses are needed to return geometry; disc and bar selection impose further tests. Acceptance remains provisional.'),
('Definitions and implementation. ', 'Photutils defines the convergence statuses and parameter errors: https://photutils.readthedocs.io/en/stable/api/photutils.isophote.Isophote.html . Our additional strict and relaxed screening rules are implemented in accepted() in Euclid/adaptive_MER_VIS/fit_adaptive_MER_VIS_geometry.py.')
]
anchor=ps[10]
for lead,body in items:
 el=OxmlElement('w:p');anchor._p.addnext(el);q=Paragraph(el,anchor._parent);q.style=d.styles['Normal'];q.add_run(lead).bold=True;q.add_run(body);anchor=q
# Keep definitions elsewhere consistent with the more precise sampling description.
for p in d.paragraphs:
 if p.text.startswith('Intensity is the fitted average brightness'):
  p.text=p.text.replace('Coverage is the usable fraction of sampled ellipse positions.', 'Coverage is the usable fraction of sampled positions along a particular ellipse, not a fraction of the galaxy area. Normal convergence means Photutils reached its numerical stopping condition; it does not validate the physical component identification.')
d.core_properties.comments='Clarified segmentation-area provenance and per-isophote acceptance thresholds on 26 October 2026. Numerical results and thresholds unchanged.'
d.save(out)
print(out.resolve());print('Images preserved:',len(d.inline_shapes))
